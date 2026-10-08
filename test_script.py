import sys
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Define paths
OUTPUT_DIR = r"D:\POC Projects\Vintage Melodies Android"
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "Vintage_Melodies_Existing_Project_and_Android_Blueprint.pptx")
IMG_DIR = r"D:\POC Projects\Vintage Melodies\public\images"

BG_IMG = os.path.join(IMG_DIR, "background.jpg")
WEB_IMG = os.path.join(IMG_DIR, "web.png")
MOBILE_IMG = os.path.join(IMG_DIR, "Mobile.png")

# Initialize presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette
DARK_BG = RGBColor(0x18, 0x10, 0x08)       # Deep Vintage Dark Coffee
CARD_BG = RGBColor(0x2A, 0x1C, 0x10)       # Translucent Warm Brown Card
CARD_BORDER = RGBColor(0x6E, 0x48, 0x24)   # Subtle Amber Border
AMBER_ACCENT = RGBColor(0xF6, 0x82, 0x1F)  # Bright Vintage Amber
GOLD_ACCENT = RGBColor(0xE6, 0x9D, 0x45)   # Warm Gold Accent
TEXT_WHITE = RGBColor(0xFF, 0xF8, 0xF0)    # Soft Off-White
TEXT_MUTED = RGBColor(0xC8, 0xB2, 0x9B)    # Muted Tan Text
HEADER_GOLD = RGBColor(0xFF, 0xD1, 0x80)   # Highlight Gold
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
TABLE_HEADER_BG = RGBColor(0x40, 0x25, 0x12)
TABLE_ROW_BG = RGBColor(0x24, 0x18, 0x0C)

blank_layout = prs.slide_layouts[6]

def set_slide_background(slide):
    # Set dark background fill for slide
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = DARK_BG
    bg_shape.line.fill.background()
    return bg_shape

def add_header(slide, title_text, category_text="VINTAGE MELODIES ARCHITECTURE BLUEPRINT"):
    # Header bar
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = AMBER_ACCENT
    
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE

def add_card(slide, left, top, width, height, title="", border_color=CARD_BORDER, bg_color=CARD_BG):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.5)
    
    if title:
        tb = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.15), Inches(width - 0.4), Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = GOLD_ACCENT
    return card

def add_bullet_list(slide, left, top, width, height, items, font_size=13, space_after=8):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            header, desc = item
            run1 = p.add_run()
            run1.text = "• " + header + ": "
            run1.font.bold = True
            run1.font.size = Pt(font_size)
            run1.font.color.rgb = HEADER_GOLD
            
            run2 = p.add_run()
            run2.text = desc
            run2.font.size = Pt(font_size)
            run2.font.color.rgb = TEXT_WHITE
        else:
            run = p.add_run()
            run.text = "• " + item
            run.font.size = Pt(font_size)
            run.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(space_after)

print("Helper functions initialized.")
